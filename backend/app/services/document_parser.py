"""Parse PDF, DOCX, and TXT files into the canonical Document model."""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import fitz
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph as DocxParagraph

from app.core.exceptions import DocumentParseError, UnsupportedFileTypeError
from app.core.logging import get_logger
from app.models.document import Document, Paragraph, Section
from app.services.file_storage import SUPPORTED_EXTENSIONS

logger = get_logger(__name__)

HEADING_STYLE_PREFIXES = ("heading", "title")


class DocumentParser:
    """Format-specific parsers that preserve order, pages, and headings."""

    def parse(
        self,
        file_path: str | Path,
        *,
        original_filename: str | None = None,
        document_id: str | None = None,
    ) -> Document:
        """Parse a file from disk into a Document. Never raises for a single bad block."""
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError()
        if not path.exists():
            raise DocumentParseError(f"File not found: {path.name}")

        filename = original_filename or path.name
        doc_id = document_id or str(uuid4())

        try:
            if extension == ".pdf":
                document = self._parse_pdf(path, filename, doc_id)
            elif extension == ".docx":
                document = self._parse_docx(path, filename, doc_id)
            else:
                document = self._parse_txt(path, filename, doc_id)
        except UnsupportedFileTypeError:
            raise
        except DocumentParseError:
            raise
        except Exception:
            logger.exception("Failed to parse document filename=%s", filename)
            raise DocumentParseError("The document could not be parsed.") from None

        document.stored_path = str(path)
        logger.info(
            "document parsed filename=%s type=%s pages=%s paragraphs=%s sections=%s",
            document.filename,
            document.file_type,
            document.total_pages,
            len(document.paragraphs),
            len(document.sections),
        )
        return document

    def _parse_pdf(self, path: Path, filename: str, document_id: str) -> Document:
        try:
            pdf = fitz.open(path)
        except Exception:
            logger.exception("Malformed PDF")
            raise DocumentParseError("The PDF file is malformed or unreadable.") from None

        sections: list[Section] = []
        paragraphs: list[Paragraph] = []
        current_section: Section | None = None
        page_count = pdf.page_count
        avg_size = self._estimate_pdf_font_size(pdf)

        try:
            for page_number, page in enumerate(pdf, start=1):
                try:
                    blocks = page.get_text("dict").get("blocks", [])
                except Exception:
                    logger.warning("Skipping unreadable PDF page %s", page_number)
                    continue

                for block in blocks:
                    try:
                        if block.get("type") != 0:
                            continue
                        text, max_size = self._pdf_block_text(block)
                        if not text:
                            continue
                        if self._looks_like_heading(text, font_size=max_size, average_size=avg_size):
                            current_section = self._start_section(sections, text)
                            continue
                        paragraph = self._make_paragraph(
                            index=len(paragraphs) + 1,
                            text=text,
                            section=current_section,
                            page_number=page_number,
                        )
                        paragraphs.append(paragraph)
                        if current_section is not None:
                            current_section.paragraph_ids.append(paragraph.id)
                    except Exception:
                        logger.warning("Skipping malformed PDF paragraph on page %s", page_number)
                        continue
        finally:
            pdf.close()

        return Document(
            id=document_id,
            filename=filename,
            file_type="pdf",
            total_pages=page_count,
            sections=sections,
            paragraphs=paragraphs,
        )

    def _estimate_pdf_font_size(self, pdf: fitz.Document) -> float:
        sizes: list[float] = []
        for page in pdf:
            try:
                for block in page.get_text("dict").get("blocks", []):
                    if block.get("type") != 0:
                        continue
                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            size = float(span.get("size") or 0)
                            if size > 0:
                                sizes.append(size)
            except Exception:
                continue
            if len(sizes) > 400:
                break
        if not sizes:
            return 12.0
        return sum(sizes) / len(sizes)

    def _pdf_block_text(self, block: dict) -> tuple[str, float]:
        lines: list[str] = []
        max_size = 0.0
        for line in block.get("lines", []):
            parts: list[str] = []
            for span in line.get("spans", []):
                parts.append(span.get("text") or "")
                max_size = max(max_size, float(span.get("size") or 0))
            line_text = "".join(parts).strip()
            if line_text:
                lines.append(line_text)
        return "\n".join(lines).strip(), max_size

    def _parse_docx(self, path: Path, filename: str, document_id: str) -> Document:
        try:
            source = DocxDocument(str(path))
        except Exception:
            logger.exception("Malformed DOCX")
            raise DocumentParseError("The DOCX file is malformed or unreadable.") from None

        sections: list[Section] = []
        paragraphs: list[Paragraph] = []
        current_section: Section | None = None

        for block in source.element.body.iterchildren():
            try:
                if block.tag == qn("w:p"):
                    paragraph = DocxParagraph(block, source)
                    text = (paragraph.text or "").strip()
                    if not text:
                        continue
                    if self._is_docx_heading(paragraph, text):
                        current_section = self._start_section(sections, text)
                        continue
                    item = self._make_paragraph(
                        index=len(paragraphs) + 1,
                        text=text,
                        section=current_section,
                    )
                    paragraphs.append(item)
                    if current_section is not None:
                        current_section.paragraph_ids.append(item.id)
                elif block.tag == qn("w:tbl"):
                    table = Table(block, source)
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if not cells:
                            continue
                        text = " | ".join(cells)
                        item = self._make_paragraph(
                            index=len(paragraphs) + 1,
                            text=text,
                            section=current_section,
                        )
                        paragraphs.append(item)
                        if current_section is not None:
                            current_section.paragraph_ids.append(item.id)
            except Exception:
                logger.warning("Skipping malformed DOCX block")
                continue

        return Document(
            id=document_id,
            filename=filename,
            file_type="docx",
            total_pages=1,
            sections=sections,
            paragraphs=paragraphs,
        )

    def _parse_txt(self, path: Path, filename: str, document_id: str) -> Document:
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            logger.exception("Malformed TXT")
            raise DocumentParseError("The text file could not be read.") from None

        sections: list[Section] = []
        paragraphs: list[Paragraph] = []
        current_section: Section | None = None

        raw_blocks = text.replace("\r\n", "\n").split("\n\n")
        for block in raw_blocks:
            try:
                cleaned = "\n".join(line.rstrip() for line in block.split("\n")).strip()
                if not cleaned:
                    continue
                first_line, *rest = cleaned.split("\n")
                if self._looks_like_heading(first_line) and (
                    not rest or self._looks_like_heading(cleaned)
                ):
                    current_section = self._start_section(sections, first_line)
                    remainder = "\n".join(rest).strip()
                    if not remainder:
                        continue
                    cleaned = remainder
                item = self._make_paragraph(
                    index=len(paragraphs) + 1,
                    text=cleaned,
                    section=current_section,
                    page_number=1,
                )
                paragraphs.append(item)
                if current_section is not None:
                    current_section.paragraph_ids.append(item.id)
            except Exception:
                logger.warning("Skipping malformed TXT paragraph")
                continue

        return Document(
            id=document_id,
            filename=filename,
            file_type="txt",
            total_pages=1,
            sections=sections,
            paragraphs=paragraphs,
        )

    def _start_section(self, sections: list[Section], title: str) -> Section:
        section = Section(
            id=f"sec-{len(sections) + 1:04d}",
            title=" ".join(title.replace("#", " ").split()),
        )
        sections.append(section)
        return section

    def _make_paragraph(
        self,
        *,
        index: int,
        text: str,
        section: Section | None,
        page_number: int | None = None,
    ) -> Paragraph:
        return Paragraph(
            id=f"p-{index:04d}",
            text=text,
            section_id=section.id if section else None,
            section_title=section.title if section else None,
            page_number=page_number,
        )

    def _is_docx_heading(self, paragraph: DocxParagraph, text: str) -> bool:
        style_name = ""
        try:
            if paragraph.style and paragraph.style.name:
                style_name = paragraph.style.name.lower()
        except Exception:
            style_name = ""
        if any(style_name.startswith(prefix) for prefix in HEADING_STYLE_PREFIXES):
            return True
        return self._looks_like_heading(text)

    def _looks_like_heading(
        self,
        text: str,
        *,
        font_size: float | None = None,
        average_size: float | None = None,
    ) -> bool:
        compact = " ".join(text.split())
        if not compact or len(compact) > 80:
            return False
        if compact.startswith("#"):
            return True
        if compact.isupper() and any(ch.isalpha() for ch in compact) and len(compact.split()) <= 10:
            return True
        if font_size and average_size and font_size >= average_size * 1.35 and len(compact.split()) <= 12:
            return True
        return False
