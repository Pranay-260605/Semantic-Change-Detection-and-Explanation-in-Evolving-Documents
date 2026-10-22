from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from app.core.exceptions import DocumentParseError, UnsupportedFileTypeError
from app.services.document_parser import DocumentParser

parser = DocumentParser()


def test_txt_parsing(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
    document = parser.parse(path)
    assert document.file_type == "txt"
    assert document.total_pages == 1
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "First paragraph.",
        "Second paragraph.",
    ]


def test_txt_headings_become_sections(tmp_path: Path) -> None:
    path = tmp_path / "headed.txt"
    path.write_text("OVERVIEW\n\nThe product launched in 2026.", encoding="utf-8")
    document = parser.parse(path)
    assert document.sections
    assert document.sections[0].title == "OVERVIEW"
    assert document.paragraphs[0].section_title == "OVERVIEW"


def test_docx_parsing(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    source = DocxDocument()
    source.add_heading("Introduction", level=1)
    source.add_paragraph("Hello from Word.")
    source.save(path)
    document = parser.parse(path)
    assert document.file_type == "docx"
    assert any("Hello from Word." in paragraph.text for paragraph in document.paragraphs)
    assert any(section.title == "Introduction" for section in document.sections)


def test_pdf_parsing_preserves_pages(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Revenue was ten million.")
    pdf.save(path)
    pdf.close()
    document = parser.parse(path)
    assert document.file_type == "pdf"
    assert document.total_pages == 1
    assert document.paragraphs
    assert document.paragraphs[0].page_number == 1
    assert "Revenue" in document.paragraphs[0].text


def test_invalid_extension(tmp_path: Path) -> None:
    path = tmp_path / "notes.png"
    path.write_bytes(b"not-a-document")
    with pytest.raises(UnsupportedFileTypeError):
        parser.parse(path)


def test_malformed_pdf(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf")
    with pytest.raises(DocumentParseError):
        parser.parse(path)
